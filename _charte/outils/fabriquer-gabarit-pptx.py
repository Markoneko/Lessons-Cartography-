"""Fabrique _charte/gabarit-geodata.pptx, le modèle utilisé pour la sortie PowerPoint.

On part du modèle neutre du convertisseur (dont les noms de mises en page sont
imposés) et on y applique la charte Géodata : couleurs, polices, logo, strates.
À relancer seulement si la charte change :

    quarto pandoc -o /tmp/neutre.pptx --print-default-data-file reference.pptx
    python3 _charte/outils/fabriquer-gabarit-pptx.py /tmp/neutre.pptx
"""
import sys, re, zipfile, shutil, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

ICI = os.path.dirname(os.path.abspath(__file__))
CHARTE = os.path.dirname(ICI)
SORTIE = os.path.join(CHARTE, "gabarit-geodata.pptx")
BLEU, ORANGE, GRIS, BLANC = RGBColor(0x28, 0x3B, 0x89), RGBColor(0xE3, 0x82, 0x26), RGBColor(0x6E, 0x6F, 0x7A), RGBColor(0xFF, 0xFF, 0xFF)
COULEURS = {"dk2": "283B89", "lt2": "F4F5FB", "accent1": "283B89", "accent2": "E38226", "accent3": "541C5B",
            "accent4": "8A5264", "accent5": "9655A2", "accent6": "D8D33A", "hlink": "283B89", "folHlink": "541C5B"}

def logo(n): return os.path.join(CHARTE, "logos", n)

def poser_image(cible, fichier, gauche, haut, largeur=None, hauteur=None):
    """Ajoute une image sur un maître ou une mise en page (non prévu par python-pptx)."""
    from pptx.oxml.shapes.picture import CT_Picture
    partie, rid = cible.part.get_or_add_image_part(fichier)
    l, h = partie.image.size
    if largeur is None: largeur = int(hauteur * l / h)
    if hauteur is None: hauteur = int(largeur * h / l)
    arbre = cible.shapes._spTree
    pic = CT_Picture.new_pic(len(arbre) + 100, os.path.basename(fichier), "", rid, gauche, haut, largeur, hauteur)
    arbre.append(pic)

def style_texte(forme, taille, couleur, gras=False, police=None, alignement=PP_ALIGN.LEFT):
    for p in forme.text_frame.paragraphs or []:
        p.alignment = alignement
    f = forme.text_frame.paragraphs[0].font
    f.size, f.bold, f.color.rgb = Pt(taille), gras, couleur
    if police: f.name = police

p = Presentation(sys.argv[1])
maitre = p.slide_masters[0]

# Maître : titre à gauche en Raleway bleu, texte en Roboto, logo en bas à gauche
for ph in maitre.placeholders:
    t = str(ph.placeholder_format.type)
    if "TITLE" in t:
        ph.left, ph.top, ph.width, ph.height = Inches(0.5), Inches(0.28), Inches(9.0), Inches(0.75)
        style_texte(ph, 26, BLEU, gras=True, police="Raleway")
    elif "BODY" in t:
        ph.left, ph.top, ph.width, ph.height = Inches(0.5), Inches(1.15), Inches(9.0), Inches(3.75)
        style_texte(ph, 16, RGBColor(0x2B, 0x2B, 0x33))
    elif "SLIDE_NUMBER" in t:
        ph.left, ph.top, ph.width = Inches(8.5), Inches(5.17), Inches(1.0)
        style_texte(ph, 10, BLEU, gras=True, alignement=PP_ALIGN.RIGHT)
    elif "FOOTER" in t:
        ph.left, ph.top, ph.width = Inches(1.6), Inches(5.17), Inches(6.5)
        style_texte(ph, 9, GRIS)
    elif "DATE" in t:
        ph.left, ph.width = Inches(8.0), Inches(0.1)
poser_image(maitre, logo("geodata.png"), Inches(0.5), Inches(5.1), hauteur=Inches(0.33))

for mep in p.slide_layouts:
    if mep.name in ("Title and Content", "Two Content", "Comparison", "Title Only"):
        for ph in mep.placeholders:
            i = ph.placeholder_format.idx
            if i == 0:
                ph.left, ph.top, ph.width, ph.height = Inches(0.5), Inches(0.28), Inches(9.0), Inches(0.75)
        if mep.name == "Title and Content":
            c = mep.placeholders[1]
            c.left, c.top, c.width, c.height = Inches(0.5), Inches(1.15), Inches(9.0), Inches(3.75)
        if mep.name == "Two Content":
            for i, gauche in ((1, 0.5), (2, 5.08)):
                c = mep.placeholders[i]
                c.left, c.top, c.width, c.height = Inches(gauche), Inches(1.15), Inches(4.42), Inches(3.75)
    elif mep.name == "Title Slide":
        for ph in mep.placeholders:
            i = ph.placeholder_format.idx
            if i == 0:
                ph.left, ph.top, ph.width, ph.height = Inches(0.5), Inches(1.7), Inches(5.3), Inches(1.4)
                style_texte(ph, 36, BLEU, gras=True, police="Raleway")
            elif i == 1:
                ph.left, ph.top, ph.width, ph.height = Inches(0.5), Inches(3.2), Inches(5.3), Inches(1.0)
                style_texte(ph, 18, ORANGE, police="Raleway")
        poser_image(mep, logo("strates-orange-haut.png"), Inches(6.0), Inches(0.55), largeur=Inches(3.4))
        poser_image(mep, logo("strates-bleu-bas.png"), Inches(6.0), Inches(4.3), largeur=Inches(3.4))
    elif mep.name == "Section Header":
        fond = mep.background.fill; fond.solid(); fond.fore_color.rgb = BLEU
        for ph in mep.placeholders:
            i = ph.placeholder_format.idx
            if i == 0:
                ph.left, ph.top, ph.width, ph.height = Inches(0.8), Inches(1.6), Inches(8.4), Inches(1.2)
                style_texte(ph, 36, BLANC, gras=True, police="Raleway")
            elif i == 1:
                ph.left, ph.top, ph.width, ph.height = Inches(0.8), Inches(2.9), Inches(8.4), Inches(0.9)
                style_texte(ph, 18, BLANC)
        poser_image(mep, logo("strates-blanc-bas.png"), Inches(0.8), Inches(4.4), largeur=Inches(4.2))

p.save(SORTIE)

# Thème : couleurs et polices de la charte (python-pptx ne sait pas le faire, on corrige le XML)
tmp = SORTIE + ".tmp"
with zipfile.ZipFile(SORTIE) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename.startswith("ppt/theme/theme") and item.filename.endswith(".xml"):
            x = data.decode("utf-8")
            for nom, hexa in COULEURS.items():
                x = re.sub(r"(<a:%s>)\s*<a:(srgbClr|sysClr)[^>]*/>\s*(</a:%s>)" % (nom, nom),
                           r'\1<a:srgbClr val="%s"/>\3' % hexa, x)
            x = re.sub(r'(<a:majorFont>\s*<a:latin typeface=")[^"]*', r"\1Raleway", x)
            x = re.sub(r'(<a:minorFont>\s*<a:latin typeface=")[^"]*', r"\1Roboto", x)
            data = x.encode("utf-8")
        if item.filename == "ppt/slideMasters/slideMaster1.xml":
            # Style de titre et tailles de texte par défaut
            x = data.decode("utf-8")
            x = re.sub(r'(<p:titleStyle><a:lvl1pPr) algn="ctr"', r'\1 algn="l"', x)
            d, f = x.index("<p:titleStyle>"), x.index("</p:titleStyle>")
            titre = re.sub(r'sz="\d+"', 'sz="2600" b="1"', x[d:f]).replace('val="tx1"', 'val="tx2"')
            x = x[:d] + titre + x[f:]
            debut, fin = x.index("<p:bodyStyle>"), x.index("</p:bodyStyle>")
            corps = x[debut:fin]
            for avant, apres in (("2400", "1600"), ("2100", "1400"), ("1800", "1300"), ("1500", "1200")):
                corps = corps.replace('sz="%s"' % avant, 'sz="_%s"' % apres)
            x = x[:debut] + corps.replace('sz="_', 'sz="') + x[fin:]
            data = x.encode("utf-8")
        if item.filename.startswith("ppt/slideLayouts/") and item.filename.endswith(".xml"):
            data = data.decode("utf-8").replace('algn="ctr"', 'algn="l"').encode("utf-8")
        zout.writestr(item, data)
shutil.move(tmp, SORTIE)
print("Écrit :", SORTIE)
